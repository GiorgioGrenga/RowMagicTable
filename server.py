"""FastAPI backend and static web server for the table converter."""

from __future__ import annotations

import io
import json
import os
import re
import zipfile
from copy import copy
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from ai_processor import (
    MAX_DOCUMENT_BYTES,
    TableExtractionError,
    _normalize_table,
    extract_table,
)

ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
MAX_WORKBOOK_BYTES = 15 * 1024 * 1024
MAX_PREVIEW_ROWS = 4
EXCEL_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

app = FastAPI(title="Tabella in Excel", docs_url=None, redoc_url=None)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/", include_in_schema=False)
def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/health", include_in_schema=False)
def health() -> dict[str, str]:
    return {"status": "ok"}


def _workbook_from_bytes(file_bytes: bytes):
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Il file Excel è vuoto.")
    if len(file_bytes) > MAX_WORKBOOK_BYTES:
        raise HTTPException(status_code=413, detail="Il file Excel supera 15 MB.")
    if not zipfile.is_zipfile(io.BytesIO(file_bytes)):
        raise HTTPException(
            status_code=400,
            detail="File non valido: carica un file Excel .xlsx non protetto.",
        )
    try:
        return load_workbook(io.BytesIO(file_bytes))
    except (InvalidFileException, KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise HTTPException(
            status_code=400,
            detail="Non riesco a leggere il file Excel. Verifica che sia un .xlsx valido.",
        ) from exc


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value).strip()


def _sheet_header(sheet) -> tuple[int, list[dict[str, Any]]]:
    header_row = next(
        (
            row_number
            for row_number in range(1, sheet.max_row + 1)
            if any(cell.value is not None for cell in sheet[row_number])
        ),
        1,
    )
    headers: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    for column_number in range(1, sheet.max_column + 1):
        value = sheet.cell(header_row, column_number).value
        if value is None:
            continue
        name = _cell_text(value) or f"Colonna {column_number}"
        seen[name] = seen.get(name, 0) + 1
        label = f"{name} ({seen[name]})" if seen[name] > 1 else name
        headers.append(
            {"label": label, "value": label, "column": column_number}
        )
    return header_row, headers


def _sheet_metadata(sheet) -> dict[str, Any]:
    header_row, headers = _sheet_header(sheet)
    preview = []
    if headers:
        for row_number in range(header_row + 1, min(sheet.max_row, header_row + MAX_PREVIEW_ROWS) + 1):
            preview.append(
                [
                    _cell_text(sheet.cell(row_number, header["column"]).value)
                    for header in headers
                ]
            )
    return {
        "name": sheet.title,
        "header_row": header_row,
        "headers": [header["label"] for header in headers],
        "preview": preview,
        "empty": not headers,
    }


def _load_table_json(raw_table: str) -> dict[str, list[Any]]:
    try:
        data = json.loads(raw_table)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="La tabella inviata non è valida.") from exc
    try:
        return _normalize_table(data)
    except TableExtractionError as exc:
        raise HTTPException(status_code=exc.http_status, detail=str(exc)) from exc


def _style_header(cell) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="204A43")
    cell.alignment = Alignment(vertical="center")


def _set_excel_value(sheet, row: int, column: int, value: Any):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def _write_new_workbook(table: dict[str, list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tabella"
    for column, header in enumerate(table["headers"], start=1):
        _style_header(_set_excel_value(sheet, 1, column, header))
    for row_number, row in enumerate(table["rows"], start=2):
        for column, value in enumerate(row, start=1):
            _set_excel_value(sheet, row_number, column, value)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    _fit_columns(sheet)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _fit_columns(sheet) -> None:
    for column_number in range(1, sheet.max_column + 1):
        longest = max(
            (
                len(_cell_text(sheet.cell(row, column_number).value))
                for row in range(1, min(sheet.max_row, 100) + 1)
            ),
            default=10,
        )
        sheet.column_dimensions[get_column_letter(column_number)].width = min(
            max(longest + 2, 12), 40
        )


def _append_to_workbook(
    workbook_bytes: bytes,
    table: dict[str, list[Any]],
    sheet_name: str,
    raw_mapping: str,
) -> bytes:
    workbook = _workbook_from_bytes(workbook_bytes)
    if sheet_name not in workbook.sheetnames:
        raise HTTPException(status_code=400, detail="Il foglio Excel selezionato non esiste.")
    try:
        mapping = json.loads(raw_mapping)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="La mappatura delle colonne non è valida.") from exc
    if not isinstance(mapping, list) or len(mapping) != len(table["headers"]):
        raise HTTPException(status_code=400, detail="Completa la mappatura di tutte le colonne.")

    sheet = workbook[sheet_name]
    header_row, existing = _sheet_header(sheet)
    existing_by_name: dict[str, list[int]] = {}
    for header in existing:
        existing_by_name.setdefault(header["value"].casefold(), []).append(header["column"])

    destinations: list[int | None] = []
    used_destinations: set[int] = set()
    added_headers: set[str] = set()
    next_column = max((header["column"] for header in existing), default=0) + 1
    for index, entry in enumerate(mapping):
        if not isinstance(entry, dict) or entry.get("mode") not in {"map", "add", "skip"}:
            raise HTTPException(status_code=400, detail="Una scelta di colonna non è valida.")
        mode = entry["mode"]
        if mode == "skip":
            destinations.append(None)
            continue
        if mode == "map":
            target = entry.get("target")
            matches = existing_by_name.get(target.casefold(), []) if isinstance(target, str) else []
            if not matches:
                raise HTTPException(
                    status_code=400,
                    detail=f"La colonna di destinazione per «{table['headers'][index]}» non esiste.",
                )
            destination = matches[0]
        else:
            name = entry.get("name")
            if not isinstance(name, str) or not name.strip():
                raise HTTPException(
                    status_code=400, detail="Inserisci un nome per ogni nuova colonna."
                )
            name = name.strip()
            folded = name.casefold()
            if folded in existing_by_name or folded in added_headers:
                raise HTTPException(
                    status_code=400,
                    detail=f"La colonna «{name}» esiste già. Abbinala alla colonna presente.",
                )
            added_headers.add(folded)
            destination = next_column
            next_column += 1
            header_cell = _set_excel_value(sheet, header_row, destination, name)
            _style_header(header_cell)
        if destination in used_destinations:
            raise HTTPException(
                status_code=400,
                detail="Due colonne estratte non possono essere abbinate alla stessa colonna Excel.",
            )
        used_destinations.add(destination)
        destinations.append(destination)

    if not used_destinations:
        raise HTTPException(status_code=400, detail="Seleziona almeno una colonna da aggiungere.")

    last_data_row = max(
        (
            row_number
            for row_number in range(1, sheet.max_row + 1)
            if any(cell.value is not None for cell in sheet[row_number])
        ),
        default=header_row,
    )
    for row_offset, source_row in enumerate(table["rows"], start=1):
        output_row = last_data_row + row_offset
        for source_index, destination in enumerate(destinations):
            if destination is not None:
                cell = _set_excel_value(
                    sheet, output_row, destination, source_row[source_index]
                )
                if last_data_row > header_row:
                    previous = sheet.cell(last_data_row, destination)
                    if previous.has_style:
                        cell._style = copy(previous._style)
                    if previous.number_format:
                        cell.number_format = previous.number_format

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _safe_filename(filename: str) -> str:
    stem = Path(filename).stem.strip() or "tabella"
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-._") or "tabella"
    return f"{safe}.xlsx"


def _excel_response(file_bytes: bytes, filename: str) -> Response:
    return Response(
        content=file_bytes,
        media_type=EXCEL_CONTENT_TYPE,
        headers={
            "Content-Disposition": (
                f"attachment; filename=\"{_safe_filename(filename)}\"; "
                f"filename*=UTF-8''{quote(_safe_filename(filename))}"
            ),
            "Cache-Control": "no-store",
        },
    )


@app.post("/api/extract")
def extract(
    document_file: UploadFile = File(...),
    workbook_file: UploadFile | None = File(default=None),
) -> dict[str, Any]:
    document_bytes = document_file.file.read(MAX_DOCUMENT_BYTES + 1)
    workbook_info = None
    if workbook_file is not None and workbook_file.filename:
        if not workbook_file.filename.lower().endswith(".xlsx"):
            raise HTTPException(
                status_code=400, detail="Per concatenare i dati carica un file .xlsx."
            )
        workbook = _workbook_from_bytes(
            workbook_file.file.read(MAX_WORKBOOK_BYTES + 1)
        )
        workbook_info = {
            "sheets": [_sheet_metadata(sheet) for sheet in workbook.worksheets]
        }
    try:
        table = extract_table(
            document_bytes,
            document_file.filename or "",
            os.getenv("GEMINI_API_KEY"),
        )
    except TableExtractionError as exc:
        raise HTTPException(status_code=exc.http_status, detail=str(exc)) from exc

    return {"table": table, "workbook": workbook_info}


@app.post("/api/export")
def export(
    table_json: str = Form(...),
    filename: str = Form(default="tabella.xlsx"),
    workbook_file: UploadFile | None = File(default=None),
    sheet_name: str = Form(default=""),
    mapping_json: str = Form(default="[]"),
) -> Response:
    table = _load_table_json(table_json)
    if workbook_file is None or not workbook_file.filename:
        return _excel_response(_write_new_workbook(table), filename)

    if not workbook_file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Il file deve essere in formato .xlsx.")
    workbook_bytes = workbook_file.file.read(MAX_WORKBOOK_BYTES + 1)
    result = _append_to_workbook(
        workbook_bytes,
        table,
        sheet_name,
        mapping_json,
    )
    return _excel_response(result, filename)
